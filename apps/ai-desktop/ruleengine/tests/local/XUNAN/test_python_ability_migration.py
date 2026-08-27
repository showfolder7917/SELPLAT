"""当前用户能力统一迁移到 Python 的回归测试。"""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest


PROJECT_ROOT = next(
    candidate for candidate in Path(__file__).resolve().parents
    if (candidate / "settings.gradle").is_file()
)
RULEENGINE_ROOT = PROJECT_ROOT / "apps/ai-desktop/ruleengine"
XUNAN_CODE_ROOT = RULEENGINE_ROOT / "python/local/XUNAN"
JAVA_LOCAL_ROOT = RULEENGINE_ROOT / "java/local"
OPTION_TEMP_ROOT = PROJECT_ROOT / "OPTION/temp"
if str(XUNAN_CODE_ROOT) not in sys.path:
    sys.path.insert(0, str(XUNAN_CODE_ROOT))


def load_module(name: str, path: Path):
    """从正式能力路径加载模块，测试不复制生产实现。"""

    specification = importlib.util.spec_from_file_location(name, path)
    assert specification is not None
    assert specification.loader is not None
    module = importlib.util.module_from_spec(specification)
    sys.modules[name] = module
    specification.loader.exec_module(module)
    return module


class PythonAbilityMigrationTests(unittest.TestCase):
    """覆盖语言收敛、规则引用和关键等价行为。"""

    @classmethod
    def setUpClass(cls) -> None:
        OPTION_TEMP_ROOT.mkdir(parents=True, exist_ok=True)
        cls.pinyin = load_module(
            "pinyin_docx_tools_migration_test",
            XUNAN_CODE_ROOT / "abilities/pinyin_docx_tools.py",
        )
        cls.teaching = load_module(
            "teaching_image_tools_migration_test",
            XUNAN_CODE_ROOT / "abilities/teaching_image_tools.py",
        )
        cls.fujitsu = load_module(
            "fujitsu_excel_tools_migration_test",
            XUNAN_CODE_ROOT / "abilities/fujitsu_excel_tools.py",
        )
        cls.quality = load_module(
            "presentation_quality_inspector_migration_test",
            XUNAN_CODE_ROOT / "abilities/presentation_quality_inspector.py",
        )
        cls.pptx_tools = sys.modules["selplat_xunan_pptx_tools"]

    def test_expected_python_abilities_exist(self) -> None:
        """迁移后仍有效的旧 Java 和 Node 职责必须有正式 Python 替代入口。"""

        expected = {
            "cross_platform_tools.py",
            "pinyin_docx_tools.py",
            "teaching_image_tools.py",
            "fujitsu_excel_tools.py",
            "idiom_ppt_generator.py",
            "oral_performance_ppt_tools.py",
            "presentation_quality_inspector.py",
        }
        existing = {path.name for path in (XUNAN_CODE_ROOT / "abilities").glob("*.py")}
        self.assertTrue(expected <= existing)

    def test_xunan_has_no_java_node_or_executable_archive(self) -> None:
        """迁移完成后当前用户不得残留 Java、Node 或封存可执行代码。"""

        forbidden: list[Path] = []
        for language, suffixes in (("java", {".java"}), ("node", {".js", ".mjs", ".cjs", ".ts"})):
            root = RULEENGINE_ROOT / language / "local/XUNAN"
            if root.exists():
                forbidden.extend(path for path in root.rglob("*") if path.suffix.lower() in suffixes)
        archive_root = XUNAN_CODE_ROOT / "archive"
        if archive_root.exists():
            forbidden.extend(path for path in archive_root.rglob("*.py"))
        self.assertEqual([], forbidden)

    def test_python_source_is_only_abilities_and_util(self) -> None:
        """当前用户 Python 正式源码只允许能力与共享工具两种职责目录。"""

        misplaced = [
            path for path in XUNAN_CODE_ROOT.rglob("*.py")
            if path.relative_to(XUNAN_CODE_ROOT).parts[0] not in {"abilities", "util"}
        ]
        self.assertEqual([], misplaced)

    def test_rule_engine_is_on_demand_python_without_gradle_or_http(self) -> None:
        """rule-engine 只保留按需 Python 能力，不得恢复后端层级、旧文档或常驻 HTTP 入口。"""

        self.assertFalse((RULEENGINE_ROOT / "build.gradle").exists())
        self.assertFalse((RULEENGINE_ROOT / "backend").exists())
        self.assertFalse((RULEENGINE_ROOT / "src").exists())
        self.assertFalse((RULEENGINE_ROOT / "docs").exists())
        self.assertFalse((XUNAN_CODE_ROOT / "abilities/rule_engine_backend.py").exists())
        settings_text = (PROJECT_ROOT / "settings.gradle").read_text(encoding="utf-8")
        self.assertNotIn("include('apps:rule-engine:backend')", settings_text)
        self.assertNotIn("project(':apps:rule-engine:backend')", settings_text)
        root_build_text = (PROJECT_ROOT / "build.gradle").read_text(encoding="utf-8")
        self.assertIn("def selplatNonGradleApplicationScopes = ['rule-engine']", root_build_text)
        self.assertIn("selplatSpecialGateTaskPaths['rule-engine']", root_build_text)
        self.assertNotIn("nonJavaLeafProjectPaths", root_build_text)

    def test_vscode_has_no_rule_engine_runtime_configuration(self) -> None:
        """VS Code 不得把按需 Python 能力重新暴露为独立 Java 或 HTTP 服务。"""

        # 编辑器配置不是运行必需品；文件不存在即没有违规入口，不应制造虚假失败。
        tasks_path = PROJECT_ROOT / ".vscode/tasks.json"
        tasks = (
            json.loads(tasks_path.read_text(encoding="utf-8"))
            if tasks_path.is_file() else {"tasks": []}
        )
        labels = {task["label"] for task in tasks.get("tasks", [])}
        self.assertNotIn("acode-java: compile", labels)
        self.assertNotIn("rule-engine:classes", labels)
        self.assertNotIn("rule-engine:run", labels)
        self.assertNotIn(
            "rule_engine_backend.py",
            json.dumps(tasks, ensure_ascii=False),
        )

        launch_path = PROJECT_ROOT / ".vscode/launch.json"
        launch = (
            json.loads(launch_path.read_text(encoding="utf-8"))
            if launch_path.is_file() else {"configurations": []}
        )
        self.assertNotIn(
            "rule-engine-backend",
            {
                configuration.get("projectName")
                for configuration in launch.get("configurations", [])
            },
        )

        settings_path = PROJECT_ROOT / ".vscode/settings.json"
        if settings_path.is_file():
            settings = json.loads(settings_path.read_text(encoding="utf-8"))
            self.assertEqual(
                "${workspaceFolder}/cache/gradle-user-home",
                settings["java.import.gradle.user.home"],
            )

    def test_rule_engine_java_local_root_is_removed(self) -> None:
        """Python-only 后端不得残留未接入构建的 Java local 源码根。"""

        self.assertFalse(JAVA_LOCAL_ROOT.exists())

    def test_overridden_pinyin_cell_is_not_changed_by_sandhi(self) -> None:
        """人工纠音只保护实际命中的位置，不影响同文其他“一、不”。"""

        converter = self.pinyin.PinyinConverter()
        cells = [
            self.pinyin.PinyinCell("yī", "一"),
            self.pinyin.PinyinCell("xīn", "心"),
            self.pinyin.PinyinCell("yī", "一"),
            self.pinyin.PinyinCell("qù", "去"),
        ]
        result = converter._apply_reading_sandhi(cells, {0, 1})
        self.assertEqual("yī", result[0].pinyin)
        self.assertEqual("yí", result[2].pinyin)

    def test_invalid_override_and_source_overwrite_are_blocked(self) -> None:
        """纠音列数不一致和源目标相同必须在写文件前阻断。"""

        with tempfile.TemporaryDirectory(prefix="pinyin_migration_", dir=OPTION_TEMP_ROOT) as directory:
            root = Path(directory)
            dictionary = root / "dictionary.tsv"
            dictionary.write_text("中国\tzhōng\n", encoding="utf-8")
            with self.assertRaises(ValueError):
                self.pinyin.load_overrides(dictionary)
            source = root / "source.docx"
            source.touch()
            with self.assertRaises(ValueError):
                self.pinyin.generate(source, source)

    def test_teaching_file_name_is_cross_platform_safe(self) -> None:
        """教学图片文件名不得含三平台不稳定字符。"""

        self.assertEqual("古诗_第一课", self.teaching.sanitize_file_name("古诗:第一课"))

    def test_fujitsu_csv_export_creates_openable_workbook(self) -> None:
        """Fujitsu Python 能力必须能把 UTF-8 CSV 生成可重新打开的工作簿。"""

        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory(prefix="fujitsu_migration_", dir=OPTION_TEMP_ROOT) as directory:
            root = Path(directory)
            (root / "输入.csv").write_text("编号,名称\n1,契约\n", encoding="utf-8-sig")
            output = root / "结果.xlsx"
            result = self.fujitsu.export_database(root, output)
            self.assertEqual(1, result["sheets"])
            workbook = load_workbook(output, read_only=True)
            self.assertEqual("契约", workbook["输入"]["B2"].value)

    def test_horizontal_ppt_generation_and_inspection(self) -> None:
        """共享 PPT util 生成的横版文件必须通过 Python 通用质检。"""

        with tempfile.TemporaryDirectory(prefix="ppt_migration_", dir=OPTION_TEMP_ROOT) as directory:
            output = Path(directory) / "样例.pptx"
            presentation = self.pptx_tools.create_presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            self.pptx_tools.add_text(slide, "迁移验证", (80, 80, 400, 100), name="CONTENT_TITLE")
            presentation.save(output)
            result = self.quality.inspect_horizontal_deck(
                output,
                expected_slides=1,
                required_shape_names=("CONTENT_TITLE",),
            )
            self.assertEqual("passed", result["status"], result["errors"])

    def test_active_rules_do_not_reference_removed_xunan_sources(self) -> None:
        """当前规则正文和文档不得继续调用已删除的 XUNAN Java/Node 能力。"""

        rule_root = RULEENGINE_ROOT / "rules/local/XUNAN"
        violations = []
        for path in rule_root.rglob("*.md"):
            if "archive" in path.parts:
                continue
            text = path.read_text(encoding="utf-8")
            if "src/main/java/com/sp/selplat/local/code/XUNAN" in text:
                violations.append(path)
            if "src/main/node/com/sp/selplat/local/code/XUNAN" in text:
                violations.append(path)
        self.assertEqual([], sorted(set(violations)))

    def test_python_dependency_manifest_covers_migrated_formats(self) -> None:
        """迁移所需 PPT、Excel、DOCX、图片、XML 和拼音库必须统一登记。"""

        requirements = (RULEENGINE_ROOT / "requirements-python.txt").read_text(encoding="utf-8")
        for dependency in ("python-pptx", "openpyxl", "python-docx", "Pillow", "lxml", "pypinyin"):
            self.assertIn(dependency, requirements)


if __name__ == "__main__":
    unittest.main()
